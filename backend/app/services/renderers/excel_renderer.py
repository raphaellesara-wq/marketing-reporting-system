"""
Excel Renderer — openpyxl branded Excel workbook.
Sheets: Summary (KPI cards) | Platform Breakdown | Daily Data + Chart
"""

from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.chart.label import DataLabelList
from openpyxl.styles import (
    Alignment, Border, Fill, Font, GradientFill,
    PatternFill, Side,
)
from openpyxl.utils import get_column_letter

from app.services.report_generator import ReportData


def _hex(color: str) -> str:
    return color.lstrip("#").upper()


class ExcelRenderer:

    def __init__(self, data: ReportData):
        self.data = data
        self.primary = _hex(data.brand_primary)
        self.secondary = _hex(data.brand_secondary)
        self.wb = Workbook()

    def render(self, output_path: str):
        self.wb.remove(self.wb.active)
        self._sheet_summary()
        self._sheet_platforms()
        self._sheet_daily()
        self.wb.save(output_path)

    # ── Helpers ───────────────────────────────────────────────────────────────

    def _header_font(self, size=11):
        return Font(name="Calibri", bold=True, color="FFFFFF", size=size)

    def _header_fill(self, color=None):
        return PatternFill("solid", fgColor=color or self.primary)

    def _border(self):
        thin = Side(style="thin", color="D1D5DB")
        return Border(left=thin, right=thin, top=thin, bottom=thin)

    def _center(self):
        return Alignment(horizontal="center", vertical="center", wrap_text=True)

    def _set_col_widths(self, ws, widths: list):
        for i, w in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(i)].width = w

    def _write_title_row(self, ws, title: str, colspan: int):
        ws.append([title])
        ws.merge_cells(start_row=ws.max_row, start_column=1,
                       end_row=ws.max_row, end_column=colspan)
        cell = ws.cell(row=ws.max_row, column=1)
        cell.font = Font(name="Calibri", bold=True, color="FFFFFF", size=14)
        cell.fill = self._header_fill(self.secondary)
        cell.alignment = self._center()
        ws.row_dimensions[ws.max_row].height = 32

    def _write_header_row(self, ws, headers: list):
        ws.append(headers)
        for col, _ in enumerate(headers, 1):
            cell = ws.cell(row=ws.max_row, column=col)
            cell.font  = self._header_font()
            cell.fill  = self._header_fill()
            cell.alignment = self._center()
            cell.border = self._border()
        ws.row_dimensions[ws.max_row].height = 22

    def _write_data_row(self, ws, values: list, alt: bool = False):
        ws.append(values)
        row = ws.max_row
        bg = "F9FAFB" if alt else "FFFFFF"
        for col, _ in enumerate(values, 1):
            cell = ws.cell(row=row, column=col)
            cell.fill      = PatternFill("solid", fgColor=bg)
            cell.alignment = self._center()
            cell.border    = self._border()
            cell.font      = Font(name="Calibri", size=10)

    # ── Sheet: Summary ────────────────────────────────────────────────────────

    def _sheet_summary(self):
        ws = self.wb.create_sheet("Summary")
        d = self.data
        k = d.kpis

        self._write_title_row(ws, d.title, 4)
        ws.append([])
        ws.append([f"Client: {d.client_name}", "", f"Period: {d.date_from.strftime('%Y-%m-%d')} → {d.date_to.strftime('%Y-%m-%d')}", ""])
        ws.merge_cells(start_row=ws.max_row, start_column=1, end_row=ws.max_row, end_column=2)
        ws.merge_cells(start_row=ws.max_row, start_column=3, end_row=ws.max_row, end_column=4)
        ws.append([])

        # KPI cards (2×4 grid)
        kpi_pairs = [
            ("Total Spend",    f"${k.total_spend:,.2f}"),
            ("Impressions",    f"{k.total_impressions:,.0f}"),
            ("Clicks",         f"{k.total_clicks:,.0f}"),
            ("Conversions",    f"{k.total_conversions:,.0f}"),
            ("Avg CTR",        f"{k.avg_ctr:.2f}%"),
            ("Avg CPC",        f"${k.avg_cpc:.2f}"),
            ("Avg CPM",        f"${k.avg_cpm:.2f}"),
            ("Avg ROAS",       f"{k.avg_roas:.2f}x"),
        ]

        for i in range(0, 8, 4):
            chunk = kpi_pairs[i:i+4]
            labels = [c[0] for c in chunk]
            values = [c[1] for c in chunk]

            ws.append(labels)
            label_row = ws.max_row
            for col in range(1, 5):
                c = ws.cell(row=label_row, column=col)
                c.font = Font(name="Calibri", bold=True, size=9, color="6B7280")
                c.alignment = self._center()
                c.fill = PatternFill("solid", fgColor="F3F4F6")
            ws.row_dimensions[label_row].height = 18

            ws.append(values)
            val_row = ws.max_row
            for col in range(1, 5):
                c = ws.cell(row=val_row, column=col)
                c.font = Font(name="Calibri", bold=True, size=16, color=self.primary)
                c.alignment = self._center()
                c.fill = PatternFill("solid", fgColor="FFFFFF")
                c.border = Border(
                    bottom=Side(style="medium", color=self.primary)
                )
            ws.row_dimensions[val_row].height = 36
            ws.append([])

        self._set_col_widths(ws, [20, 20, 20, 20])

    # ── Sheet: Platform Breakdown ─────────────────────────────────────────────

    def _sheet_platforms(self):
        ws = self.wb.create_sheet("Platform Breakdown")

        self._write_title_row(ws, "Performance by Platform", 8)
        ws.append([])

        headers = ["Platform", "Spend ($)", "Impressions", "Clicks",
                   "Conversions", "CTR (%)", "CPC ($)", "ROAS"]
        self._write_header_row(ws, headers)

        for i, p in enumerate(self.data.platforms):
            self._write_data_row(ws, [
                p.platform, round(p.spend, 2), int(p.impressions),
                int(p.clicks), int(p.conversions),
                round(p.ctr, 2), round(p.cpc, 2), round(p.roas, 2),
            ], alt=i % 2 == 1)

        # Totals row
        k = self.data.kpis
        ws.append(["TOTAL",
                   round(k.total_spend, 2), int(k.total_impressions),
                   int(k.total_clicks),     int(k.total_conversions),
                   round(k.avg_ctr, 2),     round(k.avg_cpc, 2),
                   round(k.avg_roas, 2)])
        for col in range(1, 9):
            c = ws.cell(row=ws.max_row, column=col)
            c.font = Font(name="Calibri", bold=True, size=10, color="FFFFFF")
            c.fill = self._header_fill(self.secondary)
            c.alignment = self._center()
            c.border = self._border()

        self._set_col_widths(ws, [20, 14, 14, 12, 14, 10, 12, 10])

    # ── Sheet: Daily Data + Chart ─────────────────────────────────────────────

    def _sheet_daily(self):
        ws = self.wb.create_sheet("Daily Data")

        self._write_title_row(ws, "Daily Performance", 4)
        ws.append([])

        headers = ["Date", "Spend ($)", "Clicks", "Conversions"]
        self._write_header_row(ws, headers)

        data_start = ws.max_row + 1
        for i, pt in enumerate(self.data.daily):
            self._write_data_row(ws, [
                pt.date, round(pt.spend, 2),
                int(pt.clicks), int(pt.conversions),
            ], alt=i % 2 == 1)
        data_end = ws.max_row

        self._set_col_widths(ws, [14, 14, 12, 14])

        if data_end >= data_start:
            chart = BarChart()
            chart.type = "col"
            chart.title = "Daily Spend"
            chart.y_axis.title = "Spend ($)"
            chart.x_axis.title = "Date"
            chart.width = 20
            chart.height = 12
            chart.shape = 4

            spend_ref = Reference(ws, min_col=2, min_row=data_start - 1, max_row=data_end)
            cats = Reference(ws, min_col=1, min_row=data_start, max_row=data_end)
            chart.add_data(spend_ref, titles_from_data=True)
            chart.set_categories(cats)
            chart.series[0].graphicalProperties.solidFill = self.primary

            ws.add_chart(chart, f"F{data_start}")
